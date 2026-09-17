import csv
import io
from pathlib import Path

from openpyxl import Workbook

from app.ingestion.parsers import parse_file


def test_parse_txt(tmp_path: Path):
    target = tmp_path / "notes.txt"
    target.write_text("Section one.\n\nSection two.", encoding="utf-8")
    pages = parse_file("txt", target)
    assert len(pages) == 1
    assert pages[0].page_number == 1
    assert "Section two" in pages[0].text


def test_parse_csv_groups_rows(tmp_path: Path):
    target = tmp_path / "data.csv"
    with target.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["Month", "Revenue"])
        for i in range(250):
            writer.writerow([f"2026-{(i % 12) + 1:02d}", 1000 * i])
    pages = parse_file("csv", target)
    assert len(pages) == 3
    assert all("Month" in p.text for p in pages)


def test_parse_xlsx_per_sheet(tmp_path: Path):
    target = tmp_path / "sales.xlsx"
    workbook = Workbook()
    workbook.active.title = "Q1"
    workbook.active.append(["Month", "Revenue"])
    workbook.active.append(["Jan", 100])
    second = workbook.create_sheet("Q2")
    second.append(["Month", "Revenue"])
    second.append(["Apr", 400])
    workbook.save(target)
    workbook.close()
    pages = parse_file("xlsx", target)
    assert [p.page_number for p in pages] == [1, 2]


def test_parse_docx(tmp_path: Path):
    import docx

    target = tmp_path / "report.docx"
    document = docx.Document()
    document.add_paragraph("Quarterly summary paragraph one.")
    document.save(target)
    pages = parse_file("docx", target)
    assert len(pages) == 1
    assert "Quarterly summary" in pages[0].text


def test_unsupported_type(tmp_path: Path):
    try:
        parse_file("exe", tmp_path / "evil.exe")
    except ValueError as exc:
        assert "Unsupported" in str(exc)
    else:
        raise AssertionError("expected ValueError")


def test_csv_helper_imports():
    assert io is not None
