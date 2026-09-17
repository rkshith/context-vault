"""Document parsers.

Each parser returns a list of PageText(page_number, text). Page numbers are
what the user will see as citations, so they stay meaningful per format:
  - pdf  -> physical page number
  - docx -> 1 (no page concept)
  - xlsx -> worksheet index
  - csv  -> row group index
  - txt  -> 1
"""

import csv
from dataclasses import dataclass
from pathlib import Path


@dataclass
class PageText:
    page_number: int
    text: str


CSV_ROWS_PER_PAGE = 100


def parse_file(file_type: str, path: Path) -> list[PageText]:
    parsers = {
        "pdf": _parse_pdf,
        "docx": _parse_docx,
        "xlsx": _parse_xlsx,
        "csv": _parse_csv,
        "txt": _parse_txt,
    }
    parser = parsers.get(file_type)
    if parser is None:
        raise ValueError(f"Unsupported file type: {file_type}")
    return parser(path)


def _parse_pdf(path: Path) -> list[PageText]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = []
    for number, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            pages.append(PageText(number, text))
    return pages


def _parse_docx(path: Path) -> list[PageText]:
    import docx

    document = docx.Document(str(path))
    lines = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))
    text = "\n".join(lines).strip()
    return [PageText(1, text)] if text else []


def _parse_xlsx(path: Path) -> list[PageText]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    pages: list[PageText] = []
    try:
        for number, sheet in enumerate(workbook.worksheets, start=1):
            rows = []
            for row in sheet.iter_rows(values_only=True):
                cells = ["" if value is None else str(value).strip() for value in row]
                if any(cells):
                    rows.append(" | ".join(cells))
            if rows:
                pages.append(PageText(number, "\n".join(rows)))
    finally:
        workbook.close()
    return pages


def _parse_csv(path: Path) -> list[PageText]:
    pages: list[PageText] = []
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.reader(handle)
        header = next(reader, None)
        if header is None:
            return []

        def header_line() -> str:
            return " | ".join(header)

        buffer: list[str] = []
        number = 1
        for row in reader:
            buffer.append(" | ".join(row))
            if len(buffer) >= CSV_ROWS_PER_PAGE:
                pages.append(PageText(number, header_line() + "\n" + "\n".join(buffer)))
                buffer = []
                number += 1
        if buffer:
            pages.append(PageText(number, header_line() + "\n" + "\n".join(buffer)))
    return pages


def _parse_txt(path: Path) -> list[PageText]:
    text = path.read_text(encoding="utf-8", errors="replace").strip()
    return [PageText(1, text)] if text else []
