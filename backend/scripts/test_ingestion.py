"""End-to-end ingestion test run inside the backend container.

Logs in, uploads generated docx/xlsx/txt files, waits for processing,
and verifies vectors landed in Qdrant.
"""

import asyncio
import io
import time

import docx
import httpx
from openpyxl import Workbook

BASE = "http://localhost:8000/api"


def make_docx() -> bytes:
    buffer = io.BytesIO()
    document = docx.Document()
    for i in range(1, 41):
        document.add_paragraph(
            f"Paragraph {i}. The quarterly report shows steady growth. "
            f"Zebra{i} is a unique keyword embedded for retrieval testing. "
            "Revenue increased while costs remained stable across all regions."
        )
    document.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def make_xlsx() -> bytes:
    buffer = io.BytesIO()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sales"
    sheet.append(["Month", "Region", "Revenue"])
    for i in range(1, 31):
        sheet.append([f"2026-{(i % 12) + 1:02d}", f"Region{i % 4}", 1000 * i])
    workbook.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def make_txt() -> bytes:
    text = "\n\n".join(
        f"Section {i}. Local document intelligence helps teams answer questions "
        f"about their own files. Keyword pineapple{i}. Retrieval quality depends "
        "on chunking and embedding choices." for i in range(1, 21)
    )
    return text.encode("utf-8")


def main() -> None:
    client = httpx.Client(base_url=BASE, timeout=30)

    response = client.post(
        "/auth/login",
        json={"email": "alice@example.com", "password": "correct-horse-1"},
    )
    assert response.status_code == 200, response.text

    uploads = [
        ("test-report.docx", make_docx(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
        ("test-sales.xlsx", make_xlsx(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
        ("test-notes.txt", make_txt(), "text/plain"),
    ]

    ids = []
    for filename, content, mime in uploads:
        response = client.post("/documents", files={"file": (filename, content, mime)})
        print("upload", filename, "->", response.status_code)
        assert response.status_code == 202, response.text
        ids.append(response.json()["id"])

    status = {}
    for _ in range(180):
        time.sleep(1)
        for document_id in ids:
            data = client.get(f"/documents/{document_id}").json()
            status[document_id] = data
        if all(s["status"] in ("ready", "failed") for s in status.values()):
            break

    for data in status.values():
        print("final:", data["original_filename"], data["status"], "chunks:", data["chunk_count"], "pages:", data["page_count"])
    assert all(s["status"] == "ready" for s in status.values()), status

    async def qdrant_count() -> int:
        from qdrant_client import AsyncQdrantClient

        client_q = AsyncQdrantClient(url="http://qdrant:6333")
        result = await client_q.count("chunks", exact=True)
        await client_q.close()
        return result.count

    total = asyncio.run(qdrant_count())
    print("qdrant points:", total)
    assert total > 0

    print("INGESTION OK")


if __name__ == "__main__":
    main()
